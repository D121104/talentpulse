from __future__ import annotations

import hashlib
import io
import math
import os
import posixpath
import re
import selectors
import signal
import stat
import subprocess
import tempfile
import time
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from functools import lru_cache
from xml.etree import ElementTree

from docx import Document
from pypdf import PdfReader

from app.core.config import Settings
from app.core.errors import ServiceError
from app.domain.contracts import (
    MAX_EXTRACTED_TEXT_CHARS,
    MAX_PARSE_WARNING_CHARS,
    MAX_PARSE_WARNINGS,
    MAX_STRUCTURED_ITEM_CHARS,
    MAX_STRUCTURED_ITEMS,
    MediaType,
)

_OOXML_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
}
_CONTENT_TYPES_PATH = "[Content_Types].xml"
_CV_STRUCTURED_SECTIONS = ("skills", "education", "experience", "certificates")
_OCR_ALLOWED_LANGUAGES = frozenset({"eng", "vie"})
_OCR_PROCESS_ENV = {"PATH": os.defpath, "LC_ALL": "C.UTF-8"}
_OCR_READ_CHUNK_SIZE = 64 * 1024


class _OcrFailure(Exception):
    pass


class _OcrUnavailable(_OcrFailure):
    pass


class _OcrTimeout(_OcrFailure):
    pass


class _OcrLimitExceeded(_OcrFailure):
    pass


class _OcrFailed(_OcrFailure):
    pass


@dataclass(frozen=True)
class _ParseCandidate:
    extracted_text: str
    sections: dict[str, list[str]]
    warnings: list[str]


@dataclass(frozen=True)
class _OcrPagePlan:
    total_pages: int
    dimensions: tuple[tuple[float, float], ...]


@dataclass(frozen=True)
class CVParseResult:
    extracted_text: str
    content_sha256: str
    skills: list[str]
    education: list[str]
    experience: list[str]
    certificates: list[str]
    warnings: list[str]
    parser_version: str = "structured-parser-v1"


class CVParser:
    _SECTION_ALIASES: dict[str, tuple[str, ...]] = {
        "skills": (
            "skills",
            "technical skills",
            "professional skills",
            "competencies",
            "technologies",
            "hard skills",
            "soft skills",
            "tools",
            "frameworks",
            "programming",
            "kỹ năng",
            "kĩ năng",
            "kỹ năng chuyên môn",
            "kỹ năng kỹ thuật",
            "kỹ năng mềm",
            "công nghệ",
        ),
        "education": (
            "education",
            "academic",
            "academic background",
            "qualification",
            "qualifications",
            "degree",
            "university",
            "school",
            "học vấn",
            "trình độ học vấn",
            "bằng cấp",
            "đào tạo",
            "trường",
            "đại học",
            "trình độ",
            "quá trình đào tạo",
        ),
        "experience": (
            "experience",
            "work experience",
            "professional experience",
            "employment",
            "career",
            "work history",
            "kinh nghiệm",
            "kinh nghiệm làm việc",
            "kinh nghiệm nghề nghiệp",
            "quá trình làm việc",
            "lịch sử làm việc",
        ),
        "certificates": (
            "certificates",
            "certifications",
            "certification",
            "license",
            "licenses",
            "awards",
            "achievements",
            "chứng chỉ",
            "chứng nhận",
            "giấy chứng nhận",
            "giải thưởng",
            "bằng",
            "chứng nhận chuyên môn",
            "thành tích",
            "danh hiệu",
        ),
    }
    _IGNORED_SECTION_ALIASES: tuple[str, ...] = (
        "projects",
        "personal projects",
        "academic projects",
        "selected projects",
        "dự án",
        "dự án cá nhân",
    )
    _BULLET_PREFIX = re.compile(r"^(?:[-*•·▪►→◆●○■□]|\d+[.)])\s*")

    def __init__(self, settings: Settings) -> None:
        self._settings = settings

    def parse(self, content: bytes, media_type: MediaType) -> CVParseResult:
        if len(content) > self._settings.max_upload_bytes:
            raise ServiceError(
                "payload_too_large", "CV file exceeds the configured size limit.", 413
            )
        if not content:
            raise ServiceError("empty_file", "CV file is empty.")
        if media_type is MediaType.PDF:
            self._verify_pdf_signature(content)
            text = self._normalize_text(self._parse_pdf(content))
            if len(text) > self._text_limit:
                raise ServiceError(
                    "extracted_text_too_large",
                    "Extracted CV text exceeds the configured limit.",
                    413,
                )
            native = self._candidate(text)
            if self._settings.cv_ocr_enabled and self._should_attempt_ocr(native):
                return self._parse_pdf_with_optional_ocr(content, native)
            if not text:
                raise ServiceError("no_extractable_text", "The CV contains no extractable text.")
            return self._to_result(content, native, "structured-parser-v1")
        if media_type is MediaType.DOCX:
            self._inspect_docx_package(content)
            text = self._normalize_text(self._parse_docx(content))
        else:
            raise ServiceError(
                "unsupported_media_type", "Only PDF and DOCX files are supported.", 415
            )
        if not text:
            raise ServiceError("no_extractable_text", "The CV contains no extractable text.")
        if len(text) > self._text_limit:
            raise ServiceError(
                "extracted_text_too_large", "Extracted CV text exceeds the configured limit.", 413
            )
        return self._to_result(content, self._candidate(text), "structured-parser-v1")

    @property
    def _text_limit(self) -> int:
        return min(self._settings.max_extracted_chars, MAX_EXTRACTED_TEXT_CHARS)

    @staticmethod
    def _normalize_text(text: str) -> str:
        return "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()

    @classmethod
    def _candidate(cls, text: str) -> _ParseCandidate:
        sections, warnings = cls._extract_sections(text)
        return _ParseCandidate(text, sections, warnings)

    @staticmethod
    def _quality(candidate: _ParseCandidate) -> tuple[int, int]:
        section_count = sum(
            bool(candidate.sections[section]) for section in _CV_STRUCTURED_SECTIONS
        )
        item_count = sum(len(candidate.sections[section]) for section in _CV_STRUCTURED_SECTIONS)
        return section_count, item_count

    def _should_attempt_ocr(self, native: _ParseCandidate) -> bool:
        return len(native.extracted_text) < self._settings.cv_ocr_native_text_min_chars or all(
            not native.sections[section] for section in _CV_STRUCTURED_SECTIONS
        )

    @staticmethod
    def _append_warning(warnings: list[str], message: str) -> None:
        if message not in warnings and len(warnings) < MAX_PARSE_WARNINGS:
            warnings.append(message[:MAX_PARSE_WARNING_CHARS])

    def _to_result(
        self, content: bytes, candidate: _ParseCandidate, parser_version: str
    ) -> CVParseResult:
        return CVParseResult(
            extracted_text=candidate.extracted_text,
            content_sha256=hashlib.sha256(content).hexdigest(),
            skills=candidate.sections["skills"],
            education=candidate.sections["education"],
            experience=candidate.sections["experience"],
            certificates=candidate.sections["certificates"],
            warnings=candidate.warnings,
            parser_version=parser_version,
        )

    def _parse_pdf_with_optional_ocr(
        self, content: bytes, native: _ParseCandidate
    ) -> CVParseResult:
        try:
            ocr = self._ocr_pdf(content)
        except _OcrUnavailable:
            return self._handle_ocr_failure(content, native, "ocr_unavailable")
        except _OcrLimitExceeded:
            return self._handle_ocr_failure(content, native, "ocr_limit_exceeded")
        except (_OcrTimeout, _OcrFailed):
            return self._handle_ocr_failure(content, native, "ocr_failed")
        except Exception:
            return self._handle_ocr_failure(content, native, "ocr_failed")

        if self._quality(ocr) > self._quality(native):
            warnings = list(ocr.warnings)
            self._append_warning(
                warnings, "Local OCR was selected because it improved structured extraction."
            )
            selected = _ParseCandidate(ocr.extracted_text, ocr.sections, warnings)
            return self._to_result(content, selected, "structured-parser-v1+local-ocr-v1")
        if not native.extracted_text:
            raise ServiceError("no_extractable_text", "The CV contains no extractable text.")
        return self._to_result(content, native, "structured-parser-v1")

    def _handle_ocr_failure(
        self, content: bytes, native: _ParseCandidate, code: str
    ) -> CVParseResult:
        if native.extracted_text:
            warnings = list(native.warnings)
            warning = (
                "Local OCR was unavailable; native CV text was returned."
                if code == "ocr_unavailable"
                else "Local OCR could not complete; native CV text was returned."
            )
            self._append_warning(warnings, warning)
            return self._to_result(
                content,
                _ParseCandidate(native.extracted_text, native.sections, warnings),
                "structured-parser-v1",
            )
        if code == "ocr_unavailable":
            raise ServiceError(code, "Local OCR is unavailable; retry the CV parse later.", 503)
        if code == "ocr_limit_exceeded":
            raise ServiceError(
                code, "Local OCR limits prevented processing; retry the CV parse later.", 413
            )
        raise ServiceError(code, "Local OCR failed safely; retry the CV parse later.", 503)

    @classmethod
    def _normalize_heading(cls, line: str) -> str:
        heading = line.strip().rstrip(":：-–—| ").casefold()
        return re.sub(r"\s+", " ", heading)

    @classmethod
    def _detect_section(cls, line: str) -> str | None:
        heading = cls._normalize_heading(line)
        if not heading or len(heading) > 80 or cls._BULLET_PREFIX.match(line):
            return None
        for section, aliases in cls._SECTION_ALIASES.items():
            if heading in aliases:
                return section
        return None

    @classmethod
    def _detect_ignored_section(cls, line: str) -> bool:
        heading = cls._normalize_heading(line)
        if not heading or len(heading) > 80 or cls._BULLET_PREFIX.match(line):
            return False
        return heading in cls._IGNORED_SECTION_ALIASES

    @classmethod
    def _all_section_aliases(cls) -> tuple[str, ...]:
        return (
            tuple(
                alias
                for section_aliases in cls._SECTION_ALIASES.values()
                for alias in section_aliases
            )
            + cls._IGNORED_SECTION_ALIASES
        )

    @classmethod
    def _normalize_item(cls, line: str) -> str:
        item = cls._BULLET_PREFIX.sub("", line.strip())
        return re.sub(r"\s+", " ", item).strip()

    @classmethod
    @lru_cache(maxsize=1)
    def _inline_heading_pattern(cls) -> re.Pattern[str]:
        aliases = sorted(cls._all_section_aliases(), key=len, reverse=True)
        pattern = "|".join(re.escape(alias).replace(r"\ ", r"\s+") for alias in aliases)
        return re.compile(
            rf"(?<![\w-])(?P<heading>{pattern})\s*[:：]",
            flags=re.IGNORECASE,
        )

    @classmethod
    def _inline_heading_matches(cls, line: str) -> list[re.Match[str]]:
        if cls._BULLET_PREFIX.match(line):
            return []
        return list(cls._inline_heading_pattern().finditer(line))

    @classmethod
    @lru_cache(maxsize=1)
    def _flattened_heading_pattern(cls) -> re.Pattern[str]:
        aliases = sorted(cls._all_section_aliases(), key=len, reverse=True)
        pattern = "|".join(re.escape(alias).replace(r"\ ", r"\s+") for alias in aliases)
        return re.compile(
            rf"(?<![\w-])(?P<heading>{pattern})(?=\s|$)",
            flags=re.IGNORECASE,
        )

    @staticmethod
    def _is_uppercase_heading(heading: str) -> bool:
        return any(character.isalpha() for character in heading) and heading.isupper()

    @classmethod
    def _flattened_heading_content_is_structural(cls, line: str, match: re.Match[str]) -> bool:
        content = line[match.end() :].lstrip()
        if not content or cls._has_flattened_separator(content):
            return True
        first_token = re.match(r"[^\s,;|•·▪►→◆●○■□–—-]+", content)
        if first_token is None:
            return False
        token = first_token.group().casefold()
        if token in {
            "a",
            "an",
            "and",
            "are",
            "can",
            "contains",
            "include",
            "includes",
            "is",
            "matter",
            "matters",
            "mention",
            "mentions",
            "needed",
            "of",
            "required",
            "section",
            "the",
            "to",
            "with",
        }:
            return False
        return first_token.group()[0].isupper() or first_token.group()[0].isdigit()

    @staticmethod
    def _has_flattened_separator(content: str) -> bool:
        return bool(re.match(r"^\s*[|–—-•·▪►→◆●○■□]", content))

    @classmethod
    def _flattened_heading_matches(cls, line: str) -> list[re.Match[str]]:
        if cls._BULLET_PREFIX.match(line):
            return []
        matches: list[re.Match[str]] = []
        for match in cls._flattened_heading_pattern().finditer(line):
            at_line_start = match.start() == 0
            after_pipe = line[: match.start()].rstrip().endswith("|")
            is_heading_style = cls._is_uppercase_heading(match.group("heading"))
            has_structural_content = cls._flattened_heading_content_is_structural(line, match)
            if not has_structural_content:
                continue
            if at_line_start or after_pipe or (matches and is_heading_style):
                matches.append(match)
        return matches

    @classmethod
    def _flattened_content_start(cls, line: str, match: re.Match[str]) -> int:
        position = match.end()
        while position < len(line) and line[position].isspace():
            position += 1
        if position < len(line) and line[position] in "|–—-•·▪►→◆●○■□":
            position += 1
            while position < len(line) and line[position].isspace():
                position += 1
        return position

    @classmethod
    def _inline_section_items(cls, section: str, content: str) -> list[str]:
        content = re.sub(r"\s*\|\s*$", "", content).strip()
        if section == "skills":
            return re.split(r"\s*(?:[,;|•·▪►→◆●○■□])\s*", content)
        return [content]

    @classmethod
    def _extract_sections(cls, text: str) -> tuple[dict[str, list[str]], list[str]]:
        sections: dict[str, list[str]] = {name: [] for name in cls._SECTION_ALIASES}
        warnings: list[str] = []
        current_section: str | None = None
        saw_unclassified = False
        section_truncated: set[str] = set()

        def add_warning(message: str) -> None:
            if message not in warnings and len(warnings) < MAX_PARSE_WARNINGS:
                warnings.append(message[:MAX_PARSE_WARNING_CHARS])

        def store_current(lines: list[str]) -> None:
            nonlocal saw_unclassified
            if current_section is None:
                if lines:
                    saw_unclassified = True
                return
            for raw_line in lines:
                item = cls._normalize_item(raw_line)
                if not item:
                    continue
                if len(item) > MAX_STRUCTURED_ITEM_CHARS:
                    add_warning(f"An item in the {current_section} section was not parsed safely.")
                    continue
                values = sections[current_section]
                key = item.casefold()
                if any(existing.casefold() == key for existing in values):
                    continue
                if len(values) >= MAX_STRUCTURED_ITEMS:
                    section_truncated.add(current_section)
                    continue
                values.append(item)

        pending: list[str] = []
        for line in text.splitlines():
            inline_matches = cls._inline_heading_matches(line)
            if inline_matches:
                store_current(pending)
                pending = []
                prefix = line[: inline_matches[0].start()].strip()
                if prefix:
                    saw_unclassified = True
                for index, match in enumerate(inline_matches):
                    heading = match.group("heading")
                    inline_section = cls._detect_section(heading)
                    content_start = match.end()
                    content_end = (
                        inline_matches[index + 1].start()
                        if index + 1 < len(inline_matches)
                        else len(line)
                    )
                    content = line[content_start:content_end]
                    if inline_section is not None:
                        current_section = inline_section
                        store_current(cls._inline_section_items(inline_section, content))
                    elif cls._detect_ignored_section(heading):
                        current_section = None
                        if content.strip():
                            store_current([content])
                continue
            flattened_matches = cls._flattened_heading_matches(line)
            if flattened_matches:
                store_current(pending)
                pending = []
                prefix = line[: flattened_matches[0].start()].strip()
                if prefix:
                    saw_unclassified = True
                for index, match in enumerate(flattened_matches):
                    heading = match.group("heading")
                    flattened_section = cls._detect_section(heading)
                    content_start = cls._flattened_content_start(line, match)
                    content_end = (
                        flattened_matches[index + 1].start()
                        if index + 1 < len(flattened_matches)
                        else len(line)
                    )
                    content = line[content_start:content_end]
                    if flattened_section is not None:
                        current_section = flattened_section
                        store_current(cls._inline_section_items(flattened_section, content))
                    elif cls._detect_ignored_section(heading):
                        current_section = None
                        if content.strip():
                            store_current([content])
                continue
            section = cls._detect_section(line)
            if section is not None:
                store_current(pending)
                pending = []
                current_section = section
            elif cls._detect_ignored_section(line):
                store_current(pending)
                pending = []
                current_section = None
            else:
                pending.append(line)
        store_current(pending)

        if saw_unclassified:
            add_warning("Some CV content was outside recognized sections.")
        for section in sorted(section_truncated):
            add_warning(f"The {section} section was truncated to {MAX_STRUCTURED_ITEMS} items.")
        return sections, warnings

    def _ocr_pdf(self, content: bytes) -> _ParseCandidate:
        deadline = time.monotonic() + self._settings.cv_ocr_total_timeout_seconds
        plan = self._read_ocr_page_plan(content)
        if plan.total_pages == 0:
            raise _OcrLimitExceeded

        page_limit_reached = plan.total_pages > self._settings.cv_ocr_max_pages
        skipped_pixel_pages = 0
        text_parts: list[str] = []
        text_size = 0
        with tempfile.TemporaryDirectory(prefix="talentpulse-cv-ocr-") as workspace:
            try:
                os.chmod(workspace, 0o700, follow_symlinks=False)
            except (OSError, ValueError):
                raise _OcrFailed from None
            input_path = self._write_ocr_input(workspace, content)
            for page_number, (width_points, height_points) in enumerate(plan.dimensions, start=1):
                estimated_pixels = self._estimated_render_pixels(
                    width_points,
                    height_points,
                    self._settings.cv_ocr_dpi,
                    self._settings.cv_ocr_max_render_pixels,
                )
                if estimated_pixels is None:
                    skipped_pixel_pages += 1
                    continue
                rendered_path = self._render_ocr_page(input_path, workspace, page_number, deadline)
                try:
                    language_argument = self._ocr_language_argument()
                    recognition_arguments = [
                        "tesseract",
                        rendered_path,
                        "stdout",
                        "-l",
                        language_argument,
                        "--oem",
                        "1",
                        "--psm",
                        "3",
                    ]
                    recognition_timeout = self._remaining_timeout(
                        deadline, self._settings.cv_ocr_recognize_timeout_seconds
                    )
                    output = self._run_ocr_process(
                        recognition_arguments,
                        recognition_timeout,
                        self._settings.cv_ocr_max_stdout_bytes,
                    )
                finally:
                    self._unlink_temp_file(rendered_path)
                page_text = self._normalize_text(output.decode("utf-8", errors="replace"))
                if not page_text:
                    continue
                next_size = text_size + len(page_text) + (1 if text_parts else 0)
                if next_size > self._text_limit:
                    raise _OcrLimitExceeded
                text_parts.append(page_text)
                text_size = next_size

            if not text_parts:
                if skipped_pixel_pages:
                    raise _OcrLimitExceeded
                raise _OcrFailed

            text = self._normalize_text("\n".join(text_parts))
            if not text:
                raise _OcrFailed
            sections, warnings = self._extract_sections(text)
            if page_limit_reached:
                self._append_warning(
                    warnings, "Local OCR was limited to the configured page count."
                )
            if skipped_pixel_pages:
                self._append_warning(
                    warnings,
                    "Some PDF pages exceeded the local OCR pixel limit and were skipped.",
                )
            return _ParseCandidate(text, sections, warnings)

    def _read_ocr_page_plan(self, content: bytes) -> _OcrPagePlan:
        try:
            reader = PdfReader(io.BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise _OcrFailed
            page_count = len(reader.pages)
            if page_count > self._settings.max_pdf_pages:
                raise _OcrFailed
            pages_to_read = min(page_count, self._settings.cv_ocr_max_pages)
            dimensions: list[tuple[float, float]] = []
            for page_number in range(pages_to_read):
                page = reader.pages[page_number]
                width = float(page.mediabox.width)
                height = float(page.mediabox.height)
                if (
                    not math.isfinite(width)
                    or not math.isfinite(height)
                    or width <= 0
                    or height <= 0
                ):
                    raise _OcrFailed
                dimensions.append((width, height))
            return _OcrPagePlan(page_count, tuple(dimensions))
        except _OcrFailure:
            raise
        except Exception:
            raise _OcrFailed from None

    @staticmethod
    def _estimated_render_pixels(
        width_points: float, height_points: float, dpi: int, max_pixels: int
    ) -> int | None:
        try:
            width_pixels = math.ceil(width_points * dpi / 72)
            height_pixels = math.ceil(height_points * dpi / 72)
        except (OverflowError, ValueError):
            return None
        if width_pixels <= 0 or height_pixels <= 0:
            return None
        if width_pixels > max_pixels or height_pixels > max_pixels:
            return None
        if width_pixels > max_pixels // height_pixels:
            return None
        return width_pixels * height_pixels

    @staticmethod
    def _write_ocr_input(workspace: str, content: bytes) -> str:
        try:
            file_descriptor, input_path = tempfile.mkstemp(
                prefix="input-", suffix=".pdf", dir=workspace
            )
            with os.fdopen(file_descriptor, "wb") as input_file:
                input_file.write(content)
            os.chmod(input_path, 0o600, follow_symlinks=False)
            return input_path
        except (OSError, ValueError):
            if "input_path" in locals():
                CVParser._unlink_temp_file(input_path)
            raise _OcrFailed from None

    def _render_ocr_page(
        self, input_path: str, workspace: str, page_number: int, deadline: float
    ) -> str:
        output_stem = ""
        try:
            file_descriptor, output_stem = tempfile.mkstemp(prefix="page-", dir=workspace)
            os.close(file_descriptor)
            self._unlink_temp_file(output_stem)
            rendered_path = f"{output_stem}.png"
            render_arguments = [
                "pdftoppm",
                "-f",
                str(page_number),
                "-l",
                str(page_number),
                "-r",
                str(self._settings.cv_ocr_dpi),
                "-png",
                "-singlefile",
                input_path,
                output_stem,
            ]
            render_timeout = self._remaining_timeout(
                deadline, self._settings.cv_ocr_render_timeout_seconds
            )
            self._run_ocr_process(render_arguments, render_timeout, None)
            self._secure_rendered_file(rendered_path)
            return rendered_path
        except _OcrFailure:
            if output_stem:
                self._unlink_temp_file(f"{output_stem}.png")
            raise
        except (OSError, ValueError, TypeError):
            if output_stem:
                self._unlink_temp_file(f"{output_stem}.png")
            raise _OcrFailed from None

    def _ocr_language_argument(self) -> str:
        languages = tuple(self._settings.cv_ocr_languages)
        if (
            not 1 <= len(languages) <= 2
            or len(set(languages)) != len(languages)
            or any(language not in _OCR_ALLOWED_LANGUAGES for language in languages)
        ):
            raise _OcrFailed
        return "+".join(languages)

    def _remaining_timeout(self, deadline: float, configured_timeout: float) -> float:
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise _OcrTimeout
        return min(configured_timeout, remaining)

    def _run_ocr_process(
        self, arguments: Sequence[str], timeout: float, stdout_limit: int | None
    ) -> bytes:
        process: subprocess.Popen[bytes] | None = None
        try:
            process = subprocess.Popen(
                list(arguments),
                shell=False,
                close_fds=True,
                start_new_session=True,
                stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE if stdout_limit is not None else subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                env=_OCR_PROCESS_ENV.copy(),
            )
        except FileNotFoundError:
            raise _OcrUnavailable from None
        except (OSError, ValueError):
            raise _OcrFailed from None

        try:
            if stdout_limit is None:
                try:
                    return_code = process.wait(timeout=timeout)
                except subprocess.TimeoutExpired:
                    raise _OcrTimeout from None
                if return_code != 0:
                    raise _OcrFailed
                return b""

            stream = process.stdout
            if stream is None:
                raise _OcrFailed
            output = bytearray()
            deadline = time.monotonic() + timeout
            with selectors.DefaultSelector() as selector:
                selector.register(stream, selectors.EVENT_READ)
                stream_descriptor = stream.fileno()
                while True:
                    remaining = deadline - time.monotonic()
                    if remaining <= 0:
                        raise _OcrTimeout
                    if not selector.select(remaining):
                        raise _OcrTimeout
                    read_size = min(
                        _OCR_READ_CHUNK_SIZE,
                        stdout_limit + 1 - len(output),
                    )
                    if read_size <= 0:
                        raise _OcrLimitExceeded
                    chunk = os.read(stream_descriptor, read_size)
                    if not chunk:
                        break
                    output.extend(chunk)
                    if len(output) > stdout_limit:
                        raise _OcrLimitExceeded
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                raise _OcrTimeout
            try:
                return_code = process.wait(timeout=remaining)
            except subprocess.TimeoutExpired:
                raise _OcrTimeout from None
            if return_code != 0:
                raise _OcrFailed
            return bytes(output)
        except _OcrFailure:
            raise
        except (OSError, ValueError):
            raise _OcrFailed from None
        finally:
            try:
                if process.poll() is None:
                    self._terminate_ocr_process(process)
            except (OSError, ValueError):
                pass
            if process.stdout is not None:
                try:
                    process.stdout.close()
                except OSError:
                    pass

    @staticmethod
    def _terminate_ocr_process(process: subprocess.Popen[bytes]) -> None:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except (OSError, ValueError):
            try:
                process.kill()
            except (OSError, ValueError):
                return
        try:
            process.wait(timeout=1.0)
        except (OSError, subprocess.TimeoutExpired):
            pass

    @staticmethod
    def _secure_rendered_file(path: str) -> None:
        try:
            file_status = os.lstat(path)
            if stat.S_ISLNK(file_status.st_mode) or not stat.S_ISREG(file_status.st_mode):
                raise _OcrFailed
            os.chmod(path, 0o600, follow_symlinks=False)
            file_status = os.lstat(path)
            if stat.S_ISLNK(file_status.st_mode) or not stat.S_ISREG(file_status.st_mode):
                raise _OcrFailed
        except _OcrFailure:
            raise
        except (OSError, ValueError):
            raise _OcrFailed from None

    @staticmethod
    def _unlink_temp_file(path: str) -> None:
        try:
            os.unlink(path)
        except FileNotFoundError:
            pass
        except OSError:
            pass

    def _verify_pdf_signature(self, content: bytes) -> None:
        if not content.startswith(b"%PDF-"):
            raise ServiceError(
                "invalid_file_signature",
                "CV file signature does not match the declared media type.",
            )

    def _inspect_docx_package(self, content: bytes) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
                if len(infos) > self._settings.max_docx_members:
                    raise ServiceError(
                        "zip_member_count_exceeded",
                        "DOCX contains too many package members.",
                        413,
                    )
                total_uncompressed = 0
                names: set[str] = set()
                for info in infos:
                    name = info.filename
                    normalized = posixpath.normpath(name)
                    path_parts = name.split("/")
                    if (
                        not name
                        or "\x00" in name
                        or name.startswith("/")
                        or "\\" in name
                        or any(part in {".", ".."} for part in path_parts)
                        or normalized in {".", ".."}
                        or normalized.startswith("../")
                    ):
                        raise ServiceError(
                            "unsafe_zip_path",
                            "DOCX contains an unsafe package path.",
                        )
                    if normalized in names:
                        raise ServiceError(
                            "invalid_docx", "DOCX package contains duplicate members."
                        )
                    names.add(normalized)
                    if info.is_dir():
                        continue
                    if (info.external_attr >> 16) & 0o170000 == 0o120000:
                        raise ServiceError(
                            "unsafe_zip_path",
                            "DOCX contains an unsafe package entry.",
                        )
                    if info.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
                        raise ServiceError(
                            "invalid_docx", "DOCX uses an unsupported package compression method."
                        )
                    if info.file_size > self._settings.max_docx_member_uncompressed_bytes:
                        raise ServiceError(
                            "zip_member_size_exceeded",
                            "DOCX contains an oversized package member.",
                            413,
                        )
                    total_uncompressed += info.file_size
                    if total_uncompressed > self._settings.max_docx_uncompressed_bytes:
                        raise ServiceError(
                            "decompression_limit_exceeded",
                            "DOCX decompressed content exceeds the configured limit.",
                            413,
                        )
                    if name.lower().endswith((".zip", ".jar", ".7z", ".rar", ".tar", ".gz")):
                        raise ServiceError(
                            "nested_archive_rejected",
                            "DOCX contains a nested archive.",
                        )
                required = {_CONTENT_TYPES_PATH, "word/document.xml"}
                if not required.issubset(names):
                    raise ServiceError(
                        "invalid_docx", "DOCX package is missing required OOXML parts."
                    )
                try:
                    root = ElementTree.fromstring(archive.read(_CONTENT_TYPES_PATH))
                except (ElementTree.ParseError, KeyError, RuntimeError) as exc:
                    raise ServiceError(
                        "invalid_docx", "DOCX content types could not be read safely."
                    ) from exc
                content_type = (
                    "{http://schemas.openxmlformats.org/package/2006/content-types}Override"
                )
                document_types = {
                    element.attrib.get("ContentType")
                    for element in root.findall(content_type)
                    if element.attrib.get("PartName") == "/word/document.xml"
                }
                if not document_types.intersection(_OOXML_CONTENT_TYPES):
                    raise ServiceError(
                        "invalid_docx", "DOCX content types do not identify an OOXML document."
                    )
        except ServiceError:
            raise
        except (zipfile.BadZipFile, OSError, ValueError) as exc:
            raise ServiceError("invalid_docx", "DOCX could not be parsed safely.") from exc

    def _parse_pdf(self, content: bytes) -> str:
        try:
            reader = PdfReader(io.BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise ServiceError("encrypted_file", "Encrypted PDF files are not supported.")
            if len(reader.pages) > self._settings.max_pdf_pages:
                raise ServiceError(
                    "page_limit_exceeded", "PDF exceeds the configured page limit.", 413
                )
            text_parts: list[str] = []
            text_size = 0
            for page in reader.pages:
                page_text = page.extract_text() or ""
                text_size += len(page_text)
                if text_size > self._text_limit:
                    raise ServiceError(
                        "extracted_text_too_large",
                        "Extracted CV text exceeds the configured limit.",
                        413,
                    )
                text_parts.append(page_text)
            return "\n".join(text_parts)
        except ServiceError:
            raise
        except Exception as exc:
            raise ServiceError("parse_failed", "PDF could not be parsed safely.") from exc

    def _parse_docx(self, content: bytes) -> str:
        try:
            document = Document(io.BytesIO(content))
            paragraphs: list[str] = []
            text_size = 0
            for paragraph in document.paragraphs:
                text_size += len(paragraph.text)
                if text_size > self._text_limit:
                    raise ServiceError(
                        "extracted_text_too_large",
                        "Extracted CV text exceeds the configured limit.",
                        413,
                    )
                paragraphs.append(paragraph.text)
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text_size += len(cell.text)
                        if text_size > self._text_limit:
                            raise ServiceError(
                                "extracted_text_too_large",
                                "Extracted CV text exceeds the configured limit.",
                                413,
                            )
                        paragraphs.append(cell.text)
            return "\n".join(paragraphs)
        except ServiceError:
            raise
        except (zipfile.BadZipFile, ValueError) as exc:
            raise ServiceError("invalid_docx", "DOCX could not be parsed safely.") from exc
        except Exception as exc:
            raise ServiceError("parse_failed", "DOCX could not be parsed safely.") from exc
